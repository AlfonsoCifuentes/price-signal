#!/usr/bin/env python3
"""Open the workbook in Microsoft Excel (Windows, pywin32), force a full recalculation,
save the cached values, then compare every headline value with public_analysis.json.
Exit code 1 on any mismatch or any Excel error value (#N/A, #VALUE!, ...)."""
from pathlib import Path
import json, sys
import win32com.client as win32
import openpyxl

HERE = Path(__file__).resolve().parent
XLSX = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / 'PRICE_SIGNAL_Independent_Synthetic_Analysis.xlsx'  # optional: the Spanish edition
A = json.loads((HERE.parent / '02_PUBLIC_CASE_STUDY' / 'public_analysis.json').read_text(encoding='utf-8'))
h, t = A['headline'], A['traps']

xl = win32.DispatchEx('Excel.Application')
xl.Visible = False
xl.DisplayAlerts = False
try:
    wb = xl.Workbooks.Open(str(XLSX))
    xl.CalculateFull()
    wb.Save()
    wb.Close(False)
finally:
    xl.Quit()

wb = openpyxl.load_workbook(XLSX, data_only=True)
errors, checks = [], 0
for ws in wb:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('#') and c.value.rstrip('!?/0A').upper() in ('#N', '#VALUE', '#REF', '#DIV', '#NAME', '#NUM', '#NULL', '#SPILL', '#CALC'):
                errors.append(f'{ws.title}!{c.coordinate} = {c.value}')

# Summary rows 6..22 in build order; keyed by position so the Spanish edition (same rows, translated labels) is checked too
SUMMARY_ROWS = ['Product records', 'Invented series', 'Mean list price', 'Median list price', 'Min list price', 'Max list price',
                'Q1 (inclusive)', 'Q3 (inclusive)', 'Upper IQR fence', 'Comparable cross-level groups', '…with one price at every level',
                'Consistency rate', 'Exceptions', 'Stable groups after series repair', 'Consistency rate after repair',
                'With/Without Key pairs at parity', 'Mean Part A + B premium']
S = {k: wb['Summary'].cell(row=6 + i, column=3).value for i, k in enumerate(SUMMARY_ROWS)}
expect = {
    'Product records': h['products'], 'Invented series': h['series'], 'Mean list price': h['mean'],
    'Median list price': h['median'], 'Min list price': h['min'], 'Max list price': h['max'],
    'Q1 (inclusive)': t['tb_outlier']['q1'], 'Q3 (inclusive)': t['tb_outlier']['q3'],
    'Upper IQR fence': t['tb_outlier']['upper_fence'],
    'Comparable cross-level groups': h['comparables'], '…with one price at every level': h['consistent'],
    'Consistency rate': h['consistent_pct'] / 100, 'Exceptions': h['exceptions'],
    'Stable groups after series repair': h['consistent_after_series_repair'],
    'Consistency rate after repair': h['consistent_pct_after_series_repair'] / 100,
    'With/Without Key pairs at parity': t['answer_support']['parity'],
}
for k, v in expect.items():
    checks += 1
    got = S.get(k)
    tol = 0.0006 if isinstance(v, float) and v < 1 else 0.006
    if got is None or abs(float(got) - float(v)) > tol:
        errors.append(f'Summary "{k}": Excel {got!r} vs analysis {v!r}')

ex = wb['Exceptions']
for i, e in enumerate(A['exceptions'], start=6):
    checks += 1
    if abs(ex.cell(row=i, column=9).value - e['gap_eur']) > 0.005:
        errors.append(f'Exceptions row {i}: gap {ex.cell(row=i, column=9).value} vs {e["gap_eur"]}')
sv = wb['Split_Volumes']
for i, s in enumerate(t['split_volume_levels'], start=6):
    checks += 1
    if abs(sv.cell(row=i, column=7).value - s['premium_eur']) > 0.005:
        errors.append(f'Split row {i}')
qr = [wb['Quality_Register'].cell(row=r, column=4).value for r in range(6, 16)]
checks += 1
if qr[0] != 2 or qr[2] != 1 or qr[4] != 1 or qr[5] != 1:
    errors.append(f'Quality register counts unexpected: {qr}')
mismatch_rows = [c.row - 1 for c in wb['Enriched']['V'][1:] if c.value == 'MISMATCH']
checks += 1
if mismatch_rows != [m['row_id'] for m in t['series_mismatch']]:
    errors.append(f'Series_Check rows {mismatch_rows}')

print(f'Workbook verification: {checks} checks, {len(errors)} problems')
print('Summary values:', {k: (round(v, 4) if isinstance(v, float) else v) for k, v in S.items()})
print('Quality register evidence:', qr)
for e in errors:
    print('FAIL', e)
sys.exit(1 if errors else 0)
